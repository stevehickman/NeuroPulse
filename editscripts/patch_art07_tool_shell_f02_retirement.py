"""
Close OI-ART-07: retire NP-TOOL-SHELL-001 F-02 (hub port cover anchor posts) in favour of
NP-TOOL-HUB-001 F-02, and record the retirement in the document itself (Rev 2 -> Rev 3).

Why HUB owns it. SHELL Sec 1 lists "control hub tooling" as out of scope of this very
document, and its Sec 9 table says hub anchor posts need a separate document -- which became
NP-TOOL-HUB-001. HUB's <=20 mm free-length tether is the RISK-HUB-01 hazard control
(FAI-HUB-23, BLOCKING); SHELL has no reach limit. HUB is BASELINED with an issued FAI.

What does NOT transfer by default. SHELL named the right ports (USB-C + hub accessory
ports) and HUB names a DFU/service port no other document has, while omitting the accessory
ports. That is raised as NP-TOOL-HUB-001 OI-HTOOL-08, not fixed by invention here.

Convention (patch_hexzm_tool_shell_supersession.py): nothing is deleted -- F-02's
geometry, SH-06..SH-08, FAI-CV-05/06 and the F-02 rows of Sec 3/Sec 4 stay verbatim under a
banner; header status note; revision-history row.

Idempotent: re-running is a no-op.
"""

from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SHELL_PATH = "docs/np_tool_shell_001.docx"
BANNER_MARKER = "F-02 RETIRED (2026-09-23, OI-ART-07 closed)"

GREY = (0x40, 0x40, 0x40)
RED = (0xC0, 0x00, 0x00)

BLOCKS = [
    (
        BANNER_MARKER + " — hub port covers are owned by NP-TOOL-HUB-001 F-02. "
        "Do not release anything in §2.2 to a tooling manufacturer.",
        True,
        RED,
    ),
    (
        "Retired, not deleted: the §2.2 geometry below, the F-02 rows of §2 and §3, the F-02 "
        "material rows of §4 (Shore 40A TPE tether; the F-02 half of the sealant row), SH-06..SH-08, "
        "FAI-CV-05 and FAI-CV-06, and the F-02 halves of OI-01, OI-02 and OI-06 are retained verbatim "
        "as the record of what was specified.",
        False,
        GREY,
    ),
    (
        "Why NP-TOOL-HUB-001 owns them. (1) Scope: §1 of this document lists \"control hub tooling\" "
        "as NOT in scope, and §9 says hub anchor posts need a separate document — F-02 was hub "
        "geometry written into a shell specification that excluded the hub. (2) Hazard control: "
        "NP-TOOL-HUB-001's <=20 mm free-length tether maximum is the control for RISK-HUB-01 "
        "(NP-RISK-004) — a detached-but-tethered cover must not reach the hub fan intake — verified by "
        "FAI-HUB-23 (was FAI-HTOOL-02), BLOCKING. This feature has no reach limit. (3) Status: "
        "NP-TOOL-HUB-001 is BASELINED with an issued FAI checklist (NP-FAI-HUB-001); this document is "
        "releasable for F-04 only.",
        False,
        GREY,
    ),
    (
        "Magnitude correction to the 2026-08-18 banner above. The F-02 tether is a loop of 50 mm "
        "nominal CIRCUMFERENCE (§2.2), not 50 mm of free length. A loop spans at most about half its "
        "circumference between two anchors, so its reach is about 25 mm — roughly 1.25x the 20 mm "
        "maximum, not 2.5x. It still exceeds a hazard-control maximum, which is disqualifying at any "
        "ratio.",
        False,
        GREY,
    ),
    (
        "What this retirement does NOT carry over by default. This feature named the ports correctly "
        "— USB-C plus the hub accessory ports that the intranasal probe, VNS clip, cervical VNS and "
        "mastoid pad connect to — and NP-TOOL-HUB-001 F-02 does not: it names a \"DFU/service\" port no "
        "other document describes (DFU runs over USB-C) and omits the accessory ports. So the port "
        "inventory is raised as NP-TOOL-HUB-001 OI-HTOOL-08, and the tether-capture adequacy of that "
        "document's 1.0 x 0.5 mm boss (borrowed from an alignment boss) as OI-HTOOL-09. The "
        "anchor-post + through-slot pattern of §2.1/§3 (boss 4.0 mm dia x 3.5 mm, slot 2.0 x 2.5 mm, "
        "base radius >= 1.0 mm) stays in this document as reusable engineering and is the geometry "
        "F-04's dust-cover tether still refers to.",
        False,
        GREY,
    ),
]


def set_cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def already_patched(doc):
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                if BANNER_MARKER in c.text:
                    return True
    return False


def patch_docx():
    doc = Document(SHELL_PATH)
    if already_patched(doc):
        print(f"  Already patched -- skipping: {SHELL_PATH}")
        return

    anchor = None
    for p in doc.paragraphs:
        if p.text.strip().startswith("2.2  F-02"):
            anchor = p
            break
    if anchor is None:
        raise RuntimeError("Could not find the '2.2  F-02' heading")

    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.rows[0].cells[0]
    set_cell_bg(cell, "FCE4D6")
    para = cell.paragraphs[0]
    for i, (text, bold, color) in enumerate(BLOCKS):
        if i:
            para.add_run().add_break()
            para.add_run().add_break()
        r = para.add_run(text)
        r.bold = bold
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(*color)

    # Banner sits directly under the §2.2 heading, ahead of the retired text.
    anchor._p.addnext(t._tbl)

    hdr = doc.tables[0].rows[0].cells[1]
    hr = hdr.add_paragraph().add_run(
        "F-02 RETIRED 2026-09-23 — hub port covers are NP-TOOL-HUB-001 F-02 (OI-ART-07). "
        "Rev 3. Still releasable for F-04 only."
    )
    hr.bold = True
    hr.font.size = Pt(9)
    hr.font.color.rgb = RGBColor(*RED)

    rev = doc.tables[-1]
    row = rev.add_row()
    row.cells[0].text = "3"
    row.cells[1].text = "2026-09-23"
    row.cells[2].text = "NeurOne Systems Engineering"
    row.cells[3].text = (
        "OI-ART-07 closed. F-02 (accessory port cover anchor posts x3) RETIRED in favour of "
        "NP-TOOL-HUB-001 F-02, which is BASELINED, owns the hub (outside this document's §1 scope) "
        "and carries the <=20 mm tether hazard control for RISK-HUB-01 (FAI-HUB-23, BLOCKING). "
        "Nothing deleted: §2.2, SH-06..SH-08, FAI-CV-05/06 and the F-02 rows of §2-§4 retained "
        "verbatim under a banner. Corrects the Rev 2 banner's '2.5x': the 50 mm tether is a loop "
        "circumference, reach ~25 mm, ~1.25x the limit. The port inventory this feature named "
        "correctly and NP-TOOL-HUB-001 does not (hub accessory ports) is raised as OI-HTOOL-08 there; "
        "boss tether-capture adequacy as OI-HTOOL-09. F-01 and F-03 remain retired (Rev 2); F-04 "
        "unchanged. Re-scope remains OI-ART-01."
    )

    doc.save(SHELL_PATH)
    print(f"  F-02 retirement banner inserted: {SHELL_PATH}")


if __name__ == "__main__":
    patch_docx()
