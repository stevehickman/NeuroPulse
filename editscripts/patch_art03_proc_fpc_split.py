"""
Close OI-ART-03 for NP-PROC-FPC-001: split the document into the emitter-supply controls that
survive the hex-tile architecture and the tailed-FPC connector controls that do not
(Rev 7 -> Rev 8, GitHub #394).

The split is made in place, on the NP-TOOL-SHELL-001 precedent
(patch_hexzm_tool_shell_supersession.py, patch_art07_tool_shell_f02_retirement.py): nothing is
deleted. Retired sections and rows stay verbatim under a banner or a row prefix, so the record of
what was specified survives (CLAUDE.md Sec 18: retire is not delete; NP-CONV-001 Sec 4.0.4).

What survives, and why:
  Sec 2  LED emitters (V_f bin, wavelength bin, package, incoming inspection, disposition,
         candidate assessment) -- RISK-08 CARRIED to NP-RISK-003. Sec 2.1's +-0.10 V bin is
         load-bearing for NP-HW-HEXTILE-001 Sec 8.1.1 string construction (Rev 4).
  Sec 5  RA copper and the laminate rows -- RISK-10 CARRIED to NP-RISK-004; it applies to any
         flex, now the A5 L1 rigid-flex laminate and the static cluster tail, and to the A2 tile FPC.
  Sec 6  PDMS window -- RISK-04 CARRIED to NP-RISK-003; OI-HEXTILE-12 inherits the bond unchanged.
  Sec 7  the inspection rows for the above.

What is retired, and why (every one against an NP-RISK-002 Sec 3 disposition, not by judgement):
  Sec 3  BCR421W driver -- RISK-07 RETIRED: D-3 puts a driver on every tile.
  Sec 4  Hirose FH34S / JAE FF03 connector -- RISK-01 and RISK-09 RETIRED: hex tiles have no
         tail and no connector (NP-DRV-SHELL-002 Sec 8.2).
  Sec 5  connector-tail hard gold; dynamic bend radius (RISK-11 RETIRED; the dynamic set is empty,
         NP-DRV-SHELL-002 REQ-BR2-02, and the static cluster-tail bend is REQ-BR2-01); power trace
         width >= 2.0 mm and 1 oz copper, which were sized for the zone FPC's 3.6 A per zone and are
         now OI-HEXTILE-12's to set for a 1.04 A tile (NP-RISK-002 RISK-02).
  Sec 7  the inspection rows for the above.

A retired row is retired because its SUBJECT no longer exists, not because its derivation could not
be found. That distinction is NP-CONV-001 Sec 7.1's, and it is why the 1 oz / 2.0 mm rows point at
OI-HEXTILE-12 rather than simply disappearing: the tile FPC still needs a copper weight and a trace
width, and nobody has set them.

Idempotent: re-running is a no-op.
"""

from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

PATH = "docs/np_proc_fpc_001.docx"
MARKER = "SPLIT Rev 8 (2026-09-25, OI-ART-03)"
RETIRED = "⚠ RETIRED Rev 8 (OI-ART-03) — "

GREY = (0x40, 0x40, 0x40)
RED = (0xC0, 0x00, 0x00)

HEAD_BANNER = [
    (
        "⚠  " + MARKER + " — this document now has two parts. PART A (LIVE): §2 LED emitters, §5 "
        "laminate and RA copper, §6 PDMS optical window, and their §7 inspection rows. PART B "
        "(RETIRED, retained verbatim): §3 BCR421W driver, §4 Hirose FH34S / JAE FF03 connector, and "
        "the connector-tail and zone-FPC current-sizing rows of §5 and §7. Do not place a purchase "
        "order against anything marked RETIRED.",
        True,
        RED,
    ),
    (
        "Why the split falls where it does. Every retirement follows an NP-RISK-002 §3 disposition, "
        "not a judgement made here. RISK-01 and RISK-09 (connector) and RISK-07 (BCR421U/W) are "
        "RETIRED: hex tiles have no tail and no connector (NP-DRV-SHELL-002 §8.2), and D-3 puts a "
        "driver on every tile (NP-HW-HEXTILE-001). RISK-11 (dynamic flex) is RETIRED: the dynamic "
        "set is empty (NP-DRV-SHELL-002 REQ-BR2-02). RISK-08 (V_f binning) and RISK-04 (PDMS) are "
        "CARRIED to NP-RISK-003, and RISK-10 (RA copper) to NP-RISK-004, so the sections that "
        "control them stay live.",
        False,
        GREY,
    ),
    (
        "Retired is not unfounded. The §5 rows for 1 oz copper and ≥ 2.0 mm power traces were "
        "derived, from the zone FPC's 3.6 A per zone (NP-HW-FPC-001 §10.1). They are retired because "
        "that part no longer exists. The hex-tile FPC still needs a copper weight and a trace width "
        "for a 24 V / 1.04 A tile, and OI-HEXTILE-12 owns setting them. Until it does, NO copper "
        "weight or trace width is specified by this document for the tile FPC, and a PO must take "
        "them from the fabrication drawing.",
        False,
        GREY,
    ),
    (
        "Scope after Rev 8: emitter, flex-laminate and PDMS supply for artifacts A1 (hex-tile module "
        "shell, PDMS window), A2 (hex-tile FPC and element population) and A5 (L1 inner-bowl "
        "laminate), NP-ART-001 §2.1. §2.6 still SELECTS NOTHING: OI-HEXTILE-02, OI-LED-W1 and "
        "OI-LED-01 are open, so NP-FAI-HEXFPC-001 stays unwritable.",
        False,
        GREY,
    ),
]

SEC3_BANNER = (
    RETIRED + "§3 in full. BCR421W was the shared-FPC linear driver of the zone module. RISK-07 is "
    "RETIRED (NP-RISK-002 §3): NP-HW-HEXTILE-001 D-3 puts a driver on every tile, and the drive "
    "current is now bounded by REQ-TDRV-01's I_cap, a hardware reference (D-9). The generic hazard, "
    "a driver operated beyond its rating, is inside OI-HEXTILE-07. Retained verbatim below as the "
    "record of what was specified."
)
SEC4_BANNER = (
    RETIRED + "§4 in full. Hex tiles have no FPC tail and no connector: NP-DRV-SHELL-002 §8.2 "
    "replaced the tail with a back-face compression pad array, so a module swap actuates no ZIF. "
    "RISK-01 and RISK-09 are RETIRED (NP-RISK-002 §3). The socket contact array is artifact A3, and "
    "its controls are OI-HEXTILE-11 and SH2-DRC-05a. Retained verbatim below."
)
SEC5_BANNER = (
    "⚠ §5 PARTLY RETIRED Rev 8 (OI-ART-03). LIVE: RA copper (RISK-10, now the A5 laminate, the "
    "static cluster tail and the A2 tile FPC), polyimide substrate, adhesive, ENIG on signal pads, "
    "IPC class, UL 94 V-0 and the certificate of conformance, with the RA-copper half of the PO "
    "language and the RA-copper certificate check. RETIRED (marked in the table): 1 oz copper and "
    "≥ 2.0 mm power traces (zone-FPC current sizing, now OI-HEXTILE-12), hard gold on connector-tail "
    "pads (no tail), and the ≥ 25 mm dynamic bend radius (no dynamic path; the static cluster-tail "
    "bend is NP-DRV-SHELL-002 REQ-BR2-01). The PO sentences and the incoming-inspection steps for "
    "trace width, tail gold (XRF) and 20 ZIF cycles are retired with them."
)
SEC6_NOTE = (
    "Rev 8 note: the bonding-process reference below (NP-HW-FPC-001 §9) is to a superseded document. "
    "The process is inherited unchanged by the hex tile — NP-HW-HEXTILE-001 OI-HEXTILE-12 and "
    "HT-DRC-16, NP-TOOL-HEXTILE-001 §3, and NP-FAI-001 FAI-TC01/TC02 — so §6 stays LIVE."
)

# Table 11 (§5): parameter -> reason
SEC5_RETIRE = {
    "Copper weight": "sized for the zone FPC's 3.6 A per zone; OI-HEXTILE-12 sets it for the tile.",
    "Minimum power trace width": "sized for the zone FPC (NP-HW-FPC-001 §10.1); OI-HEXTILE-12.",
    "Dynamic bend radius": "no dynamic path (REQ-BR2-02); static cluster-tail bend is REQ-BR2-01.",
}
SURFACE_NOTE = (
    " [Rev 8: the ENIG half is LIVE; the connector-tail hard-gold half is RETIRED (no tail).]"
)

# Table 13 (§7): (component, item substring) -> retire
SEC7_RETIRE = [
    ("BCR421W", ""),
    ("FH34S", ""),
    ("FPC", "Power trace width"),
    ("FPC", "Hard gold"),
]


def set_cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def banner_table(doc, blocks, fill="FCE4D6"):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.rows[0].cells[0]
    set_cell_bg(cell, fill)
    para = cell.paragraphs[0]
    for i, (text, bold, color) in enumerate(blocks):
        if i:
            para.add_run().add_break()
            para.add_run().add_break()
        r = para.add_run(text)
        r.bold = bold
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(*color)
    return t


def set_para_text(p, text):
    """Replace a paragraph's text, keeping the first run's formatting."""
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""


def prefix_cell(cell, prefix):
    p = cell.paragraphs[0]
    if p.runs:
        p.runs[0].text = prefix + p.runs[0].text
    else:
        p.add_run(prefix)
    for r in p.runs[:1]:
        r.font.color.rgb = RGBColor(*RED)


def already_patched(doc):
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                if MARKER in c.text:
                    return True
    return False


def find_para(doc, startswith):
    for p in doc.paragraphs:
        if p.text.strip().startswith(startswith):
            return p
    raise RuntimeError(f"paragraph not found: {startswith!r}")


def patch():
    doc = Document(PATH)
    if already_patched(doc):
        print(f"  Already patched -- skipping: {PATH}")
        return

    # Locate everything before mutating anything, so a missing anchor leaves the file untouched.
    title = find_para(doc, "Zone Module FPC — Procurement Requirements")
    rev = find_para(doc, "Revision:  7")
    applies = find_para(doc, "Applies to:")
    riskreg = find_para(doc, "Risk register:")
    head = find_para(doc, "⚠  CORRECTED 2026-09-21 (Rev 4)")
    s3 = find_para(doc, "3.  LED DRIVER IC")
    s4 = find_para(doc, "4.  ZONE MODULE CONNECTOR")
    s5 = find_para(doc, "5.  FPC FABRICATION REQUIREMENTS")
    s6 = find_para(doc, "6.  PDMS OPTICAL WINDOW MATERIAL")
    scope_t, s5_t, s7_t, rev_t = doc.tables[0], doc.tables[11], doc.tables[13], doc.tables[14]
    assert scope_t.rows[0].cells[0].text.startswith("Component category")
    assert s5_t.rows[0].cells[2].text.startswith("Default if not stated")
    assert s7_t.rows[0].cells[1].text.startswith("Critical inspection item")
    assert rev_t.rows[0].cells[0].text == "Rev"

    # Header block.
    set_para_text(
        title,
        "Emitter Supply, Flex Laminate and Optical Window — Procurement Requirements "
        "(was: Zone Module FPC — Procurement Requirements)",
    )
    set_para_text(rev, "Revision:  8  —  2026-09-25")
    applies.runs[-1].text = (
        "Emitter, flex-laminate and PDMS supply for NP-ART-001 artifacts A1, A2 and A5 (Rev 8). "
        "Written for the Zone Module FPC (NP-HW-FPC-001 Rev C, superseded)."
    )
    riskreg.runs[-1].text = (
        "NP-RISK-003 (RISK-04, RISK-08) and NP-RISK-004 (RISK-10). RISK-01, RISK-07 and RISK-09 are "
        "RETIRED (NP-RISK-002 §3), and §3–§4 that addressed them are retired in place."
    )
    head._p.addnext(banner_table(doc, HEAD_BANNER)._tbl)

    # §1 scope table.
    for row in scope_t.rows[1:]:
        cat = row.cells[0].text
        if cat.startswith("LED driver IC") or cat.startswith("Zone module connector"):
            prefix_cell(row.cells[0], RETIRED)
        elif cat.startswith("FPC fabrication"):
            prefix_cell(row.cells[0], "⚠ PARTLY RETIRED Rev 8 — ")

    # Section banners, directly under each heading.
    s3._p.addnext(banner_table(doc, [(SEC3_BANNER, True, RED)])._tbl)
    s4._p.addnext(banner_table(doc, [(SEC4_BANNER, True, RED)])._tbl)
    s5._p.addnext(banner_table(doc, [(SEC5_BANNER, True, RED)])._tbl)
    s6._p.addnext(banner_table(doc, [(SEC6_NOTE, False, GREY)], fill="E2EFDA")._tbl)

    # §5 table rows.
    for row in s5_t.rows[1:]:
        param = row.cells[0].text
        for key, why in SEC5_RETIRE.items():
            if param.startswith(key):
                prefix_cell(row.cells[0], RETIRED)
                row.cells[1].paragraphs[-1].add_run(f" [Retired: {why}]").font.color.rgb = RGBColor(*RED)
        if param.startswith("Surface finish"):
            row.cells[1].paragraphs[-1].add_run(SURFACE_NOTE).font.color.rgb = RGBColor(*RED)

    # §7 table rows.
    for row in s7_t.rows[1:]:
        comp, item = row.cells[0].text, row.cells[1].text
        for c, sub in SEC7_RETIRE:
            if comp.startswith(c) and (not sub or item.startswith(sub)):
                prefix_cell(row.cells[0], RETIRED)
                break

    r = rev_t.add_row()
    r.cells[0].text = "8"
    r.cells[1].text = "2026-09-25"
    r.cells[2].text = "NeurOne Systems Engineering + Procurement"
    r.cells[3].text = (
        "Rev 8 — OI-ART-03: the document is SPLIT IN PLACE into the controls that survive the "
        "hex-tile architecture (Part A, live) and the tailed-FPC controls that do not (Part B, "
        "retired and retained verbatim; GitHub #394). Every retirement follows an NP-RISK-002 §3 "
        "disposition: §3 BCR421W (RISK-07 RETIRED; D-3 per-tile driver, D-9 REQ-TDRV-01), §4 Hirose "
        "FH34S / JAE FF03 (RISK-01 and RISK-09 RETIRED; no tail, NP-DRV-SHELL-002 §8.2), and in §5 "
        "the tail hard gold, the dynamic bend radius (RISK-11 RETIRED; REQ-BR2-02's dynamic set is "
        "empty) and the 1 oz / ≥ 2.0 mm zone-FPC current sizing, which OI-HEXTILE-12 now owns for "
        "the tile. §7's matching rows are marked. LIVE: §2 emitters (RISK-08), §5 RA copper and "
        "laminate (RISK-10), §6 PDMS (RISK-04). Title, scope and risk-register lines re-scoped to "
        "artifacts A1, A2 and A5. No emitter selected and no figure changed; OI-HEXTILE-02, "
        "OI-LED-W1 and OI-LED-01 remain open."
    )

    doc.save(PATH)
    print(f"  OI-ART-03 split applied: {PATH}")


if __name__ == "__main__":
    patch()
