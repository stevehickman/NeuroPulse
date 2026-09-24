"""
ECR-EMMC-002: EMMC-FS-01 read_size / prog_size 256 -> 512 in all three columns
(UHDR, SHDR, Config).  NP-FW-EMMC-001 Rev 2 -> Rev 3.

Why. EMMC-UHDR-05 (and EMMC-SHDR-03 by reference) sets the AES-XTS data unit to
512 bytes "matching the LittleFS read_size and prog_size", while EMMC-FS-01's own
table printed 256 for both.  The two clauses could not both hold, and the one
that states its reason wins: XTS encrypts a whole data unit under one tweak, so a
256-byte program into a 512-byte unit forces the encryption layer to read,
decrypt, merge, re-encrypt and rewrite the WHOLE unit -- and a power loss inside
that rewrite damages the other 256 bytes, which littlefs has already committed.
Shown by sweep in NP-SOUP-LFS-001 Rev 4 Sec 13.1.3 (a durable Map 3 record lost
on the Config instance at 256; nothing lost at 512) and decided for all three
instances by Rev 4 Sec 13.1 (logs) and Rev 5 Sec 13.8 (Config, OI-LFS-10).

What does NOT change.  Every other EMMC-FS-01 value stands.  The firmware now
implements the Config column exactly as printed (OI-LFS-10 closed, NP-SOUP-LFS-001
Rev 6 Sec 13.9): lookahead_size 64, cache_size 512, block_cycles 200, attr_max
256, name_max 64, metadata_max 4,096.  ECR-EMMC-001 (NP-FW-NVRAM-001 Sec 3.3.1.4,
the EMMC-CFG-01/-02 clause changes) is a separate ECR and is not applied here.

Convention (patch_art07_tool_shell_f02_retirement.py): nothing is deleted -- the
superseded 256 stays visible in each changed cell; banner under the table; header
revision; revision-history row.

Idempotent: re-running is a no-op.
"""

from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

EMMC_PATH = "docs/np_fw_emmc_001.docx"
BANNER_MARKER = "ECR-EMMC-002 APPLIED (2026-09-24, Rev 3)"
NEW_CELL = "512 bytes (= the XTS data unit, EMMC-UHDR-05) [ECR-EMMC-002, Rev 3; was 256 bytes]"
OLD_HEADER = "NP-FW-EMMC-001  Rev 2  |  2026-08-12  |  Pre-Storage-Code Baseline"
NEW_HEADER = "NP-FW-EMMC-001  Rev 3  |  2026-09-24  |  Pre-Storage-Code Baseline"

GREY = (0x40, 0x40, 0x40)
RED = (0xC0, 0x00, 0x00)

BLOCKS = [
    (
        BANNER_MARKER + " — read_size and prog_size are 512 bytes in all three "
        "columns, not the 256 this table printed through Rev 2.",
        True,
        RED,
    ),
    (
        "Why. EMMC-UHDR-05 sets the XTS data unit to 512 bytes \"matching the LittleFS read_size "
        "and prog_size\"; this table printed 256, so the two clauses contradicted each other. XTS "
        "encrypts a whole data unit under one tweak, so a 256-byte program into a 512-byte unit "
        "forces the encryption layer to read-modify-write the whole unit, and a power loss inside "
        "that rewrite damages the other half — bytes LittleFS has already committed and synced. "
        "That breaks the one property of the block-device contract every power-loss claim in §5 "
        "rests on: a program disturbs nothing outside itself.",
        False,
        GREY,
    ),
    (
        "Evidence. NP-SOUP-LFS-001 Rev 4 §13.1.3 and Rev 5 §13.8: under a 512-byte "
        "read-modify-write tear model, the same power-loss sweep loses committed data at 256 "
        "(a durable Map 3 journal record, on the Config instance) and loses nothing at 512. The "
        "firmware enforces 512 at compile time (np_lfs_instance.c, np_lfs_log_instance.c) and at "
        "mount (np_lfs_config_validate, np_lfs_log_config_validate), and its tests require 256 to "
        "be refused.",
        False,
        GREY,
    ),
    (
        "Consequences. cache_size is already a multiple of 512 in every column (4,096 / 4,096 / "
        "512), so lfs_init()'s relations still hold and no other value moves. Every other value "
        "in this table stands, and the firmware now implements the Config column exactly as "
        "printed (OI-LFS-10 closed, NP-SOUP-LFS-001 Rev 6 §13.9). The superseded 256 is retained "
        "in each changed cell as the record of what was specified. ECR-EMMC-001 "
        "(NP-FW-NVRAM-001 §3.3.1.4) is a separate change and is not applied by this revision.",
        False,
        GREY,
    ),
]


def set_cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def already_patched(doc):
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                if BANNER_MARKER in c.text:
                    return True
    return False


def set_text(paragraphs, text):
    """Replace the text of `paragraphs`, keeping the first run's formatting."""
    runs = [r for p in paragraphs for r in p.runs]
    if not runs:
        paragraphs[0].add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""


def find_param_table(doc):
    for t in doc.tables:
        hdr = [c.text.strip() for c in t.rows[0].cells]
        if hdr == ["Parameter", "UHDR", "SHDR", "Config"]:
            return t
    raise RuntimeError("Could not find the EMMC-FS-01 parameter table")


def patch_docx():
    doc = Document(EMMC_PATH)
    if already_patched(doc):
        print(f"  Already patched -- skipping: {EMMC_PATH}")
        return

    table = find_param_table(doc)
    changed = 0
    for row in table.rows[1:]:
        name = row.cells[0].text.strip()
        if name not in ("read_size", "prog_size"):
            continue
        for cell in row.cells[1:]:
            if cell.text.strip() != "256 bytes":
                raise RuntimeError(f"{name}: expected '256 bytes', found {cell.text!r}")
            set_text(cell.paragraphs, NEW_CELL)
            changed += 1
    if changed != 6:
        raise RuntimeError(f"expected to change 6 cells, changed {changed}")

    banner = doc.add_table(rows=1, cols=1)
    banner.style = "Table Grid"
    cell = banner.rows[0].cells[0]
    set_cell_bg(cell, "FCE4D6")
    para = cell.paragraphs[0]
    for i, (text, bold, color) in enumerate(BLOCKS):
        if i:
            para.add_run().add_break()
            para.add_run().add_break()
        r = para.add_run(text)
        r.bold = bold
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(*color)
    # Directly under the EMMC-FS-01 table, ahead of EMMC-FS-02.
    table._tbl.addnext(banner._tbl)

    header = None
    for p in doc.paragraphs[:6]:
        if p.text.strip() == OLD_HEADER:
            header = p
            break
    if header is None:
        raise RuntimeError("Could not find the Rev 2 header line")
    set_text([header], NEW_HEADER)

    rev = doc.tables[-1]
    if [c.text.strip() for c in rev.rows[0].cells] != ["Rev", "Date", "Author", "Description"]:
        raise RuntimeError("The last table is not the revision history")
    row = rev.add_row()
    row.cells[0].text = "3"
    row.cells[1].text = "2026-09-24"
    row.cells[2].text = "NeurOne Firmware Engineering"
    row.cells[3].text = (
        "ECR-EMMC-002 applied. §5.2 EMMC-FS-01: read_size and prog_size 256 -> 512 bytes in the "
        "UHDR, SHDR and Config columns, resolving the contradiction with EMMC-UHDR-05 (XTS data "
        "unit 512 bytes \"matching the LittleFS read_size and prog_size\"). A program smaller than "
        "the XTS unit forces a read-modify-write of a unit whose other half is already committed, "
        "and a power loss inside it loses committed data — shown by sweep in NP-SOUP-LFS-001 "
        "Rev 4 §13.1.3 / Rev 5 §13.8 (GitHub #382, #424). Nothing deleted: the superseded 256 is "
        "kept in each changed cell and a banner records the change under the table. No other "
        "value changes; the firmware now implements every other EMMC-FS-01 value as printed "
        "(OI-LFS-10 closed). ECR-EMMC-001 not applied here."
    )

    doc.save(EMMC_PATH)
    print(f"  ECR-EMMC-002 applied (6 cells, banner, header, revision row): {EMMC_PATH}")


if __name__ == "__main__":
    patch_docx()
