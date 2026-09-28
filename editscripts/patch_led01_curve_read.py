"""
OI-LED-01: record the curve-read for the three parts it names (NP-PROC-FPC-001 Rev 8 -> Rev 9,
GitHub #333).

The datasheets were supplied into the session and read: ams-OSRAM GH CSSRM5.24 v1.1 (2023-01-18),
ams-OSRAM SFH 4718A v1.5 (2026-01-09) and Lumileds DS190 (dated 2018-01-09 on the supplied copy).
Each IF = f(VF) trace is a vector path in the PDF. It was extracted, converted through the plot's
own axis calibration and interpolated at 120 / 150 / 180 mA. The trace was then checked against
the datasheet's tabulated V_f where one exists. That is the Rev 5 method for the Luminus family,
applied to the parts OI-LED-01 names by name.

Nothing is deleted. Every correction appends a marked Rev 9 note to the cell or paragraph it
corrects, on the NP-TOOL-SHELL-001 precedent (NP-CONV-001 Sec 4.0.4). No part is selected.

Idempotent: re-running is a no-op.
"""

from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

PATH = "docs/np_proc_fpc_001.docx"
MARKER = "OI-LED-01 — READ FOR THE THREE NAMED PARTS (Rev 9"
R9 = " [Rev 9, GitHub #333: "

GREY = (0x40, 0x40, 0x40)
RED = (0xC0, 0x00, 0x00)

BANNER = [
    (
        "⚠  " + MARKER + ", 2026-09-28, GitHub Issue #333). The datasheets were supplied and read: "
        "ams-OSRAM GH CSSRM5.24 v1.1 (2023-01-18), ams-OSRAM SFH 4718A v1.5 (2026-01-09) and "
        "Lumileds DS190 (the supplied copy is dated 2018-01-09). Each IF = f(VF) trace was "
        "extracted from the PDF's vector path, converted through the plot's own axis calibration, "
        "interpolated, and checked against the datasheet's tabulated V_f. These are datasheet "
        "figures, not measurements, and they select nothing.",
        True,
        RED,
    ),
    (
        "V_f at 120 / 150 / 180 mA. GH CSSRM5.24: 1.77 / 1.78 / 1.80 V. Self-check PASS: the trace "
        "gives 1.990 V at 700 mA against the 1.99 V typical. SFH 4718A (single pulse, 100 µs): "
        "1.39 / 1.41 / 1.43 V. Self-check PASS: 1.744 V at 1 A against the 1.75 V typical. "
        "L1IZ-0850: the trace reads 2.83 / 2.87 / 2.91 V, but self-check FAILS. It gives 3.455 V at "
        "1000 mA against Table 2's 3.2 V typical, so the plotted device sits 0.26 V above typical. "
        "Anchoring the trace's own shape (−0.582 V from 1 A to 150 mA) to the typical gives 2.62 V. "
        "Carry the bracket 2.62–2.87 V. OI-LED-04 cites a DS190 dated 2018-06-19, which may "
        "reconcile the figure and the table. It has not been read.",
        False,
        GREY,
    ),
    (
        "String length against NP-HW-HEXTILE-001 §8.1.1 (24 V rail; N·V_f ≥ 22.4 V thermal floor, "
        "≤ 24 V − V_dropout functional ceiling). GH CSSRM5.24: N = 13, 23.15 V, 3.5 % overhead. "
        "The specified 11 × 2.10 V was a design target, and 12 falls under the floor. SFH 4718A: "
        "N = 16, 22.56 V, 6.0 %, against the specified 14 × 1.60 V. L1IZ-0850: N = 9 (23.56 V) at "
        "the anchored value or N = 8 (22.98 V) at the trace value, so it is undetermined.",
        False,
        GREY,
    ),
    (
        "Binning against §2.1. GH CSSRM5.24 ships in 0.10 V forward-voltage groups E1–F2 at 700 mA "
        "(±0.05 V), which meets §2.1 as standard. Its group edges need N = 14, 13, 12 and 12 at "
        "150 mA, so a fixed N needs a single ordered group. SFH 4718A offers NO V_f groups (ordering "
        "is by brightness group only), and its typical-to-maximum spread is 0.30 V at 1 A. L1IZ-0850 "
        "bins in 0.5 V steps (C 2.0–2.5 … F 3.5–4.0 V). Neither NIR part meets §2.1's ±0.10 V "
        "without a special binning agreement (OI-LED-06).",
        False,
        GREY,
    ),
    (
        "Radiant flux at 150 mA, against NP-HW-HEXTILE-001 §4.3's 95 mW target. GH CSSRM5.24: "
        "~233 mW typical (relative 0.219 × 1068 mW at 700 mA). SFH 4718A: ~110 mW (0.165 × 665 mW). "
        "L1IZ-0850: ~162 mW (0.155 × 1050 mW). Every candidate exceeds the target. For the 660 nm "
        "primary this becomes a blocking problem: its I_F minimum is 100 mA (\"Do not use below "
        "100 mA\"), and a 45-site T1-A CH_A is still ~663 mW/cm² at 100 mA against R-4's "
        "400 mW/cm² peak ceiling. That is NP-HW-HEXTILE-001 OI-HEXTILE-29.",
        False,
        GREY,
    ),
    (
        "Datasheet facts that correct this document, each marked in place below. GH CSSRM5.24: "
        "(1) the ordering codes are -V7A2-1-1 (Q65113A3948) and -V8A2-1-1 (Q65113A4032), and "
        "\"-V8V9-1-1-700-R33\" does not appear. (2) λpeak is given as a typical only (660 nm), and "
        "λcentroid 646 / 657 / 666 nm is a single group, so there is no wavelength bin to name on "
        "a PO. (3) No L70 or Q90 figure is published. (4) The package is 3.0 × 3.0 mm ±0.1, which "
        "fits NP-HW-HEXTILE-001 §4.1's 3.80 mm lattice (OI-HEXTILE-22). SFH 4718A: (5) the half "
        "angle is 20°, a 40° beam, not 80°. (6) λcentroid is 850 nm, and 860 nm is the peak. "
        "(7) The package height is 2.66 mm, not 2.29 mm.",
        False,
        GREY,
    ),
]

# (table index, row label prefix, column index, note)
CELL_NOTES = [
    (2, "Nominal forward voltage", 1,
     "datasheet 1.78 V for the GH CSSRM5.24 primary. The 2.0–2.2 V estimate is 0.2–0.4 V high.]"),
    (2, "Nominal forward voltage", 2,
     "datasheet 1.41 V (SFH 4718A), 2.62–2.87 V (L1IZ-0850), 2.82–2.85 V (Luminus SST, Rev 5). "
     "The 1.6–1.8 V estimate fits none of the in-window candidates.]"),
    (7, "Part number", 1,
     "the datasheet (v1.1) lists only -V7A2-1-1 (Q65113A3948) and -V8A2-1-1 (Q65113A4032). "
     "\"V8V9\" is not an ordering code there. Confirm the orderable code with ams-OSRAM.]"),
    (7, "Peak wavelength", 1,
     "λpeak 660 nm is a typical only, with no peak bin. λcentroid 646 / 657 / 666 nm min / typ / "
     "max is ONE group (\"1\"). No wavelength bin exists to specify, so the lot certificate is the "
     "only control. CLAUDE.md §3 states 660–670 nm.]"),
    (7, "Package", 1,
     "confirmed 3.0 × 3.0 mm ±0.1 (3.1 max). It fits NP-HW-HEXTILE-001 §4.1's 3.80 mm lattice with "
     "0.19 mm clearance (OI-HEXTILE-22).]"),
    (7, "Vf at ~150 mA", 1,
     "DATASHEET 1.77 / 1.78 / 1.80 V at 120 / 150 / 180 mA, read off IF = f(VF) and self-checked "
     "(1.990 V at 700 mA vs 1.99 V typical). The ~1.9 V estimate was 0.12 V high.]"),
    (7, "L70 lifetime", 1,
     "the datasheet (v1.1) publishes NO L70 or Q90 figure, only \"long lifetime\". The >102,000 h "
     "Q90 and this ✓ rest on a source this document does not cite. OI-LED-05.]"),
    (7, "Vf binning", 1,
     "CONFIRMED from the datasheet. Forward Voltage Groups E1 1.80–1.90, E2 1.90–2.00, F1 "
     "2.00–2.10 and F2 2.10–2.20 V at 700 mA are standard. A single group must be ordered for a "
     "fixed string length (NP-HW-HEXTILE-001 §8.1.1).]"),
    (8, "Peak wavelength", 2,
     "λcentroid 850 nm typ, λpeak 860 nm typ (v1.5). On centroid this part is 850 nm.]"),
    (8, "Package / beam angle", 2,
     "the datasheet (v1.5) gives a half angle of 20° typ, a 40° beam (\"OSLON Black Series "
     "(850 nm) - 40°\"), NOT 80°. The height is 2.66 mm, not 2.29. 3.75 mm square is confirmed.]"),
    (8, "Vf at ~150 mA", 1,
     "IF = f(VF) (DS190 Fig. 5) reads 2.83 / 2.87 / 2.91 V at 120 / 150 / 180 mA, but FAILS its "
     "self-check: 3.455 V at 1 A against Table 2's 3.2 V typical. Anchored to typical: 2.62 V. "
     "Bracket 2.62–2.87 V, so N = 9 or 8. V_f bins are 0.5 V wide (Table 6).]"),
    (8, "Vf at ~150 mA", 2,
     "DATASHEET 1.39 / 1.41 / 1.43 V at 120 / 150 / 180 mA (single pulse, 100 µs), self-checked "
     "(1.744 V at 1 A vs 1.75 V typical). No V_f groups are offered.]"),
]

PARA_NOTES = [
    ("OI-LED-01 [",
     "READ for GH CSSRM5.24, SFH 4718A and L1IZ-0850. The results are in the banner below: 1.78 V, "
     "1.41 V and 2.62–2.87 V at 150 mA. The L1IZ trace fails its self-check, so its string length "
     "is still undetermined. Pending-decisions §13.2e(g) item 1 is discharged for the parts it "
     "names.]"),
    ("OI-LED-02 [",
     "SFH 4718A's beam is 40°, not 80° (datasheet v1.5, half angle 20°). The \"both wide-angle\" "
     "premise holds for L1IZ-0850 (150°) and not for SFH 4718A.]"),
    ("OI-LED-03 [",
     "the GH CSSRM5.24 datasheet (v1.1) has no peak bin and a single centroid group spanning "
     "646–666 nm (typical 657 nm), so there is no bin code to specify. The λ certificate is the "
     "whole control. Ask ams-OSRAM whether a tighter selection is orderable.]"),
    ("OI-LED-05 [",
     "none of the three supplied datasheets states an L70. GH CSSRM5.24 v1.1 says \"long "
     "lifetime\". DS190 says to contact sales for radiometric power maintenance. SFH 4718A v1.5 "
     "is silent.]"),
    ("OI-LED-06 [",
     "GH CSSRM5.24's standard groups are already ±50 mV, so no agreement is needed beyond naming "
     "one group. L1IZ-0850's standard bins are 0.5 V wide, and SFH 4718A has no V_f bins, so for "
     "both this agreement is the ONLY route to §2.1.]"),
]


def set_cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def banner_table(doc, blocks, fill="FCE4D6"):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.rows[0].cells[0]
    set_cell_bg(cell, fill)
    para = cell.paragraphs[0]
    for i, (text, bold, color) in enumerate(blocks):
        if i:
            para.add_run().add_break()
            para.add_run().add_break()
        r = para.add_run(text)
        r.bold = bold
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(*color)
    return t


def red_run(paragraph, text):
    r = paragraph.add_run(text)
    r.font.color.rgb = RGBColor(*RED)
    return r


def already_patched(doc):
    return any(MARKER in c.text for t in doc.tables for r in t.rows for c in r.cells)


def find_para(doc, startswith):
    for p in doc.paragraphs:
        if p.text.strip().startswith(startswith):
            return p
    raise RuntimeError(f"paragraph not found: {startswith!r}")


def find_row(table, prefix):
    for r in table.rows[1:]:
        if r.cells[0].text.startswith(prefix):
            return r
    raise RuntimeError(f"row not found: {prefix!r}")


def patch():
    doc = Document(PATH)
    if already_patched(doc):
        print(f"  Already patched -- skipping: {PATH}")
        return

    # Locate everything before mutating anything, so a missing anchor leaves the file untouched.
    rev = find_para(doc, "Revision:  8")
    led01 = find_para(doc, "OI-LED-01 [")
    t = doc.tables
    assert t[2].rows[0].cells[1].text.startswith("660 nm LED")
    assert "GH CSSRM5.24" in t[7].rows[0].cells[1].text
    assert "L1IZ-0850" in t[8].rows[0].cells[1].text and "SFH 4718A" in t[8].rows[0].cells[2].text
    rev_t = next(x for x in t if x.rows[0].cells[0].text == "Rev" and len(x.columns) == 4)
    assert rev_t.rows[-1].cells[0].text == "8"
    cells = [(find_row(t[ti], label).cells[col], note) for ti, label, col, note in CELL_NOTES]
    paras = [(find_para(doc, start), note) for start, note in PARA_NOTES]

    rev.runs[0].text = "Revision:  9  —  2026-09-28"
    for p in rev.runs[1:]:
        p.text = ""
    for cell, note in cells:
        red_run(cell.paragraphs[-1], R9 + note)
    for para, note in paras:
        red_run(para, R9 + note)
    led01._p.addnext(banner_table(doc, BANNER)._tbl)

    r = rev_t.add_row()
    r.cells[0].text = "9"
    r.cells[1].text = "2026-09-28"
    r.cells[2].text = "NeurOne Hardware Engineering"
    r.cells[3].text = (
        "Rev 9 — OI-LED-01 READ for the three parts it names, off their own datasheets (GitHub "
        "#333). V_f at 150 mA: GH CSSRM5.24 1.78 V (self-check PASS), SFH 4718A 1.41 V (PASS), "
        "L1IZ-0850 2.62–2.87 V (trace FAILS its self-check by +0.26 V at 1 A). String lengths on the "
        "24 V rail are 13, 16 and 9-or-8, against the specified 11 and 14. Binning: GH ±0.05 V "
        "standard; SFH 4718A has no V_f groups; L1IZ bins are 0.5 V. Every candidate's 150 mA flux "
        "exceeds the 95 mW target, and the 660 nm primary cannot reach R-4 at 45 sites above its "
        "100 mA rated minimum (NP-HW-HEXTILE-001 OI-HEXTILE-29). Corrected in place: GH ordering "
        "code, wavelength binning and L70 source; SFH 4718A beam angle (40°, not 80°), centroid "
        "(850 nm) and height. No part selected; OI-LED-W1 and OI-HEXTILE-02 remain open."
    )

    doc.save(PATH)
    print(f"  OI-LED-01 curve-read recorded: {PATH}")


if __name__ == "__main__":
    patch()
